import math
from datetime import date, datetime
from io import BytesIO

import pandas as pd
import geopandas as gpd

from flask import Flask, render_template, request, jsonify, send_file
from pathlib import Path

from excel_service import (
    get_sheet_names,
    get_header_candidates,
    read_sheet_with_header,
    get_column_profile,
    get_header_values,
)

from normalization_service import (
    normalize_value,
    normalize_variety,
    normalize_for_filter,
)

from shapefile_service import (
    extract_shapefile,
    read_shapefile,
    shapefile_to_geojson,
)

from matching_service import (
    match_data,
    compare_matched_attributes,
)

app = Flask(__name__)

export_progress = {
    "active": False,
    "current": 0,
    "total": 0,
    "percent": 0,
    "status": "idle",
    "message": ""
}

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DATA_DIR.mkdir(exist_ok=True)

SHAPE_DIR = DATA_DIR / "shapefile"
SHAPE_DIR.mkdir(
    parents=True,
    exist_ok=True
)

ALLOWED_EXTENSIONS = {".xlsx", ".xls", ".xlsm"}


def get_current_excel():
    files = [
        p for p in DATA_DIR.iterdir()
        if p.is_file()
        and p.suffix.lower() in ALLOWED_EXTENSIONS
    ]

    if not files:
        return None

    return max(
        files,
        key=lambda p: p.stat().st_mtime
    )
    
# =========================================================
# FILTER EXPORT PROGRESS
# =========================================================

filter_export_progress = {
    "progress": 0,
    "status": "Menunggu...",
    "done": False
} 
    
@app.route("/")
def index():
    return render_template("index.html")

@app.route("/shape-editor")
def shape_editor():
    return render_template("shape_editor.html")


def clean_json_value(value):
    """
    Membersihkan nilai yang tidak dapat dikirim sebagai JSON valid.

    Contoh:
    NaN      -> None
    Infinity -> None
    Timestamp/date/datetime -> string ISO
    """

    if value is None:
        return None

    # Dictionary
    if isinstance(value, dict):
        return {
            key: clean_json_value(val)
            for key, val in value.items()
        }

    # List / tuple
    if isinstance(value, (list, tuple)):
        return [
            clean_json_value(val)
            for val in value
        ]

    # Float NaN / Infinity
    if isinstance(value, float):
        if not math.isfinite(value):
            return None

        return value

    # Numpy scalar seperti numpy.int64 / numpy.float64
    if hasattr(value, "item"):
        try:
            value = value.item()

            if isinstance(value, float):
                if not math.isfinite(value):
                    return None

            return value

        except (ValueError, TypeError):
            pass

    # Date / datetime
    if isinstance(value, (datetime, date)):
        return value.isoformat()

    return value

@app.route("/api/upload", methods=["POST"])
def upload_excel():

    if "file" not in request.files:
        return jsonify({
            "success": False,
            "message": "File Excel tidak ditemukan."
        }), 400

    file = request.files["file"]

    if not file.filename:
        return jsonify({
            "success": False,
            "message": "Silakan pilih file Excel."
        }), 400

    extension = Path(
        file.filename
    ).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        return jsonify({
            "success": False,
            "message": "Format harus .xlsx, .xls, atau .xlsm."
        }), 400

    # =================================================
    # HAPUS EXCEL LAMA
    # =================================================

    for old_file in DATA_DIR.iterdir():

        if (
            old_file.is_file()
            and old_file.suffix.lower() in ALLOWED_EXTENSIONS
        ):

            old_file.unlink()


    # =================================================
    # SIMPAN EXCEL BARU
    # =================================================

    destination = DATA_DIR / file.filename


    file.save(destination)

    try:

        sheets = get_sheet_names(
            destination
        )

        return jsonify({
            "success": True,
            "filename": file.filename,
            "sheets": sheets
        })

    except Exception as error:

        if destination.exists():
            destination.unlink()

        return jsonify({
            "success": False,
            "message": str(error)
        }), 500


@app.route("/api/sheets")
def sheets():

    filepath = get_current_excel()

    if filepath is None:
        return jsonify({
            "success": False,
            "message": "Belum ada file Excel."
        }), 404

    try:

        return jsonify({
            "success": True,
            "filename": filepath.name,
            "sheets": get_sheet_names(filepath)
        })

    except Exception as error:

        return jsonify({
            "success": False,
            "message": str(error)
        }), 500


@app.route("/api/headers")
def headers():

    filepath = get_current_excel()

    sheet_name = request.args.get(
        "sheet"
    )

    if filepath is None:
        return jsonify({
            "success": False,
            "message": "Belum ada file Excel."
        }), 404

    if not sheet_name:
        return jsonify({
            "success": False,
            "message": "Sheet belum dipilih."
        }), 400

    try:

        candidates = get_header_candidates(
            filepath,
            sheet_name
        )

        return jsonify({
            "success": True,
            "sheet": sheet_name,
            "candidates": candidates
        })

    except Exception as error:

        return jsonify({
            "success": False,
            "message": str(error)
        }), 500


@app.route("/api/preview")
def preview():

    filepath = get_current_excel()

    sheet_name = request.args.get(
        "sheet"
    )

    header_row = request.args.get(
        "header",
        type=int
    )

    limit = request.args.get(
        "limit",
        default=100,
        type=int
    )

    if filepath is None:
        return jsonify({
            "success": False,
            "message": "Belum ada file Excel."
        }), 404

    if not sheet_name:
        return jsonify({
            "success": False,
            "message": "Sheet belum dipilih."
        }), 400

    if not header_row or header_row < 1:
        return jsonify({
            "success": False,
            "message": "Header row tidak valid."
        }), 400

    try:

        df = read_sheet_with_header(
            filepath,
            sheet_name,
            header_row
        )
            
        preview_df = df.head(
            max(
                1,
                min(limit, 200)
            )
        ).copy()

        preview_df = preview_df.where(
            preview_df.notna(),
            None
        )

        rows = preview_df.to_dict(
            orient="records"
        )

        rows = clean_json_value(rows)

        profile = get_column_profile(df)

        profile = clean_json_value(profile)

        return jsonify({
            "success": True,
            "sheet": sheet_name,
            "header_row": header_row,
            "columns": [
                str(column)
                for column in df.columns
            ],
            "rows": rows,
            "total_rows": len(df),
            "total_columns": len(df.columns),
            "profile": profile
        })

    except Exception as error:

        return jsonify({
            "success": False,
            "message": str(error)
        }), 500


@app.route("/api/parameters", methods=["GET"])
def api_parameters():
    try:
        sheet_name = request.args.get("sheet")
        header_row = request.args.get("header", type=int)

        if not sheet_name:
            return jsonify({
                "success": False,
                "message": "Sheet belum dipilih."
            }), 400

        if header_row is None:
            return jsonify({
                "success": False,
                "message": "Header belum dipilih."
            }), 400

        filepath = get_current_excel()

        if filepath is None:
            return jsonify({
                "success": False,
                "message": "Belum ada file Excel."
            }), 404

        df = read_sheet_with_header(
            filepath,
            sheet_name,
            header_row
        )

        parameters = []

        for column in df.columns:
            series = df[column].dropna()

            if len(series) == 0:
                data_type = "empty"
            else:
                data_type = str(series.dtype)

                if data_type.startswith("int") or data_type.startswith("float"):
                    data_type = "numeric"
                elif "datetime" in data_type:
                    data_type = "date"
                else:
                    data_type = "text"

            parameters.append({
                "name": str(column),
                "type": data_type,
                "non_empty": int(series.notna().sum())
            })

        return jsonify({
            "success": True,
            "sheet": sheet_name,
            "header": header_row,
            "parameters": parameters
        })

    except Exception as e:
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500
        
        
@app.route("/api/group-values", methods=["GET"])
def api_group_values():
    try:
        sheet_name = request.args.get("sheet")
        header_row = request.args.get("header", type=int)
        parameter = request.args.get("parameter")

        if not sheet_name:
            return jsonify({
                "success": False,
                "message": "Sheet belum dipilih."
            }), 400

        if header_row is None:
            return jsonify({
                "success": False,
                "message": "Header belum dipilih."
            }), 400

        if not parameter:
            return jsonify({
                "success": False,
                "message": "Parameter belum dipilih."
            }), 400

        filepath = get_current_excel()

        if filepath is None:
            return jsonify({
                "success": False,
                "message": "Belum ada file Excel."
            }), 404

        df = read_sheet_with_header(
            filepath,
            sheet_name,
            header_row
        )

        if parameter not in df.columns:
            return jsonify({
                "success": False,
                "message": f"Kolom '{parameter}' tidak ditemukan."
            }), 404

        series = df[parameter].dropna()

        groups = {}

        for value in series:
            original = str(value).strip()

            if not original:
                continue

            # Normalisasi fleksibel
            normalization_mode = request.args.get(
                "normalization",
                "default"
            ).lower()

            normalized = normalize_for_filter(
                original,
                parameter_type="text",
                mode=normalization_mode
            )
            if normalized not in groups:
                groups[normalized] = {
                    "normalized": normalized,
                    "values": [],
                    "count": 0
                }

            if original not in groups[normalized]["values"]:
                groups[normalized]["values"].append(original)

            groups[normalized]["count"] += 1

        result = []

        for item in groups.values():
            result.append({
                "values": item["values"],
                "normalized": item["normalized"],
                "count": item["count"]
            })

        result.sort(
            key=lambda x: x["count"],
            reverse=True
        )

        return jsonify({
            "success": True,
            "parameter": parameter,
            "total_unique_groups": len(result),
            "groups": result
        })

    except Exception as e:
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500
  

@app.route("/api/filter-export-progress", methods=["GET"])
def filter_export_progress_status():

    return jsonify({
        "success": True,
        "progress": filter_export_progress["progress"],
        "status": filter_export_progress["status"],
        "done": filter_export_progress["done"]
    })

# =====================================================
# FILTER DATA
# =====================================================           

@app.route("/api/filter", methods=["POST"])
def filter_data():

    try:

        # -------------------------------------------------
        # REQUEST
        # -------------------------------------------------

        payload = request.get_json(
            silent=True
        ) or {}
        
        # -------------------------------------------------
        # RESET PROGRESS EXPORT
        # -------------------------------------------------

        if payload.get("export") is True:

            filter_export_progress["progress"] = 0
            filter_export_progress["status"] = "processing"
            filter_export_progress["done"] = False        

        sheet_name = payload.get(
            "sheet"
        )

        header_row = payload.get(
            "header"
        )

        filters = payload.get(
            "filters",
            []
        )

        # None = tampilkan seluruh hasil
        limit = payload.get(
            "limit"
        )


        # -------------------------------------------------
        # VALIDASI REQUEST
        # -------------------------------------------------

        if not sheet_name:

            return jsonify({
                "success": False,
                "message": "Sheet belum dipilih."
            }), 400


        if not header_row:

            return jsonify({
                "success": False,
                "message": "Header belum dipilih."
            }), 400


        if not isinstance(
            filters,
            list
        ):

            return jsonify({
                "success": False,
                "message": "Format filters tidak valid."
            }), 400


        # -------------------------------------------------
        # VALIDASI LIMIT
        # -------------------------------------------------

        if limit is not None:

            try:

                limit = int(
                    limit
                )

            except (
                ValueError,
                TypeError
            ):

                limit = None


            if (
                limit is not None
                and limit <= 0
            ):

                limit = None


        # -------------------------------------------------
        # FILE EXCEL
        # -------------------------------------------------

        filepath = get_current_excel()

        if not filepath:

            return jsonify({
                "success": False,
                "message": "Belum ada file Excel."
            }), 400


        # -------------------------------------------------
        # BACA SHEET
        # -------------------------------------------------

        df = read_sheet_with_header(
            filepath,
            sheet_name,
            int(header_row)
        )


        if df is None:

            return jsonify({
                "success": False,
                "message": "Data Excel tidak dapat dibaca."
            }), 400


        # -------------------------------------------------
        # RAPIKAN NAMA KOLOM
        # -------------------------------------------------

        df.columns = [
            str(column).strip()
            for column in df.columns
        ]
        
        if payload.get("export") is True:

            filter_export_progress["progress"] = 25
            filter_export_progress["status"] = "Menyiapkan parameter..."

        # -------------------------------------------------
        # DATA KOSONG
        # -------------------------------------------------

        if df.empty:

            return jsonify({
                "success": True,
                "columns": [
                    str(column)
                    for column in df.columns
                ],
                "rows": [],
                "total": 0,
                "shown": 0,
                "limit": limit,
                "show_all": limit is None
            })


        # -------------------------------------------------
        # MASK AWAL
        # -------------------------------------------------

        mask = pd.Series(
            True,
            index=df.index
        )


        # =================================================
        # PROSES SETIAP PARAMETER
        # =================================================
        #
        # Struktur baru:
        #
        # Parameter 1
        #   ├── Filter 1
        #   ├── Filter 2
        #   └── Filter 3
        #
        # Semua filter dalam satu parameter = AND
        #
        # Antar parameter diproses berdasarkan urutan:
        # Parameter 1 → Parameter 2 → Parameter 3
        #
        # Data yang sudah masuk parameter sebelumnya
        # tidak boleh masuk parameter berikutnya.
        # =================================================

        assigned_mask = pd.Series(
            False,
            index=df.index
        )
        
        comment_series = pd.Series("", index=df.index, dtype="object")


        for group_index, group in enumerate(filters, start=1):

            if not isinstance(
                group,
                dict
            ):

                continue


            conditions = group.get(
                "conditions",
                []
            )


            if not isinstance(
                conditions,
                list
            ) or not conditions:

                return jsonify({
                    "success": False,
                    "message": (
                        "Parameter tidak memiliki "
                        "filter/condition."
                    )
                }), 400


            # =================================================
            # MASK UNTUK SATU PARAMETER
            # =================================================

            group_mask = pd.Series(
                True,
                index=df.index
            )


            for condition in conditions:

                if not isinstance(
                    condition,
                    dict
                ):

                    continue


                parameter = str(
                    condition.get(
                        "parameter",
                        ""
                    )
                ).strip()


                parameter_type = str(
                    condition.get(
                        "type",
                        "text"
                    )
                ).lower()


                operator = str(
                    condition.get(
                        "operator",
                        "="
                    )
                ).upper()


                value = condition.get(
                    "value"
                )


                normalization = str(
                    condition.get(
                        "normalization",
                        "default"
                    )
                ).lower()


                # =================================================
                # VALIDASI PARAMETER
                # =================================================

                if not parameter:

                    return jsonify({
                        "success": False,
                        "message": (
                            "Ada filter yang belum "
                            "memilih atribut."
                        )
                    }), 400


                if parameter not in df.columns:

                    return jsonify({
                        "success": False,
                        "message": (
                            f"Parameter '{parameter}' "
                            f"tidak ditemukan."
                        )
                    }), 400


                series = df[
                    parameter
                ]


                # =================================================
                # TEXT / OBJECT / STRING
                # =================================================

                if parameter_type in {
                    "text",
                    "object",
                    "string"
                }:

                    # ---------------------------------------------
                    # VALUE HARUS LIST
                    # ---------------------------------------------

                    if not isinstance(
                        value,
                        list
                    ):

                        value = [
                            value
                        ]


                    # ---------------------------------------------
                    # NORMALISASI NILAI FILTER
                    # ---------------------------------------------

                    normalized_values = []


                    for item in value:

                        normalized = (
                            normalize_for_filter(
                                item,
                                parameter_type="text",
                                mode=normalization
                            )
                        )


                        if normalized != "":

                            normalized_values.append(
                                normalized
                            )


                    # ---------------------------------------------
                    # TIDAK ADA NILAI
                    # ---------------------------------------------

                    if not normalized_values:

                        return jsonify({
                            "success": False,
                            "message": (
                                f"Filter '{parameter}' "
                                "belum memiliki nilai."
                            )
                        }), 400


                    # ---------------------------------------------
                    # NORMALISASI KOLOM EXCEL
                    # ---------------------------------------------

                    normalized_series = (
                        series.apply(
                            lambda item:
                            normalize_for_filter(
                                item,
                                parameter_type="text",
                                mode=normalization
                            )
                        )
                    )


                    # ---------------------------------------------
                    # = / IN
                    # ---------------------------------------------

                    if operator in {
                        "=",
                        "IN"
                    }:

                        rule_mask = normalized_series.isin(
                            normalized_values
                        )

                    # ---------------------------------------------
                    # != / NOT IN
                    # ---------------------------------------------

                    elif operator in {
                        "!=",
                        "NOT IN"
                    }:

                        rule_mask = ~normalized_series.isin(
                            normalized_values
                        )   

                    else:

                        return jsonify({
                            "success": False,
                            "message": (
                                f"Operator '{operator}' "
                                f"tidak valid untuk "
                                f"parameter teks "
                                f"'{parameter}'."
                            )
                        }), 400


                # =================================================
                # NUMERIC
                # =================================================

                elif parameter_type == "numeric":
                   
                    numeric_series = pd.to_numeric(
                        series,
                        errors="coerce"
                    )
                

                    # ---------------------------------------------
                    # VALUE MENJADI LIST
                    # ---------------------------------------------

                    if isinstance(
                        value,
                        list
                    ):
                        numeric_values = value      
                    

                    else:

                        numeric_values = [
                            value
                        ]


                    # ---------------------------------------------
                    # KONVERSI KE FLOAT
                    # ---------------------------------------------

                    try:

                        numeric_values = [
                            float(item)
                            for item in numeric_values
                        ]

                    except (
                        ValueError,
                        TypeError
                    ):

                        return jsonify({
                            "success": False,
                            "message": (
                                f"Nilai parameter "
                                f"'{parameter}' tidak valid."
                            )
                        }), 400


                    # =================================================
                    # BETWEEN
                    # =================================================

                    if operator == "BETWEEN":

                        if len(
                            numeric_values
                        ) < 2:

                            return jsonify({
                                "success": False,
                                "message": (
                                    f"Operator BETWEEN "
                                    f"untuk '{parameter}' "
                                    "membutuhkan 2 nilai."
                                )
                            }), 400


                        minimum = numeric_values[0]

                        maximum = numeric_values[1]


                        if minimum > maximum:

                            return jsonify({
                                "success": False,
                                "message": (
                                    f"Nilai minimum "
                                    f"tidak boleh lebih besar "
                                    f"dari maksimum untuk "
                                    f"'{parameter}'."
                                )
                            }), 400


                        rule_mask = (
                            numeric_series >= minimum
                        ) & (
                            numeric_series <= maximum
                        )


                    # =================================================
                    # OPERATOR SINGLE VALUE
                    # =================================================

                    else:

                        if not numeric_values:

                            return jsonify({
                                "success": False,
                                "message": (
                                    f"Parameter '{parameter}' "
                                    "belum memiliki nilai."
                                )
                            }), 400


                        number = numeric_values[0]


                        if operator == "=":

                            rule_mask = numeric_series == number


                        elif operator == "!=":

                            rule_mask = numeric_series != number


                        elif operator == ">":

                            rule_mask = numeric_series > number


                        elif operator == ">=":

                            rule_mask = numeric_series >= number


                        elif operator == "<":

                            rule_mask = numeric_series < number


                        elif operator == "<=":

                            rule_mask = numeric_series <= number


                        else:

                            return jsonify({
                                "success": False,
                                "message": (
                                    f"Operator '{operator}' "
                                    f"tidak valid untuk "
                                    f"numeric '{parameter}'."
                                )
                            }), 400


                # =================================================
                # TYPE TIDAK DIKENAL
                # =================================================

                else:

                    return jsonify({
                        "success": False,
                        "message": (
                            f"Tipe parameter "
                            f"'{parameter_type}' "
                            f"tidak dikenal untuk "
                            f"'{parameter}'."
                        )
                    }), 400


                # =================================================
                # AND DALAM SATU PARAMETER
                # =================================================
               
                group_mask = (
                    group_mask
                    & rule_mask
                )

            # =================================================
            # HANYA DATA YANG BELUM TERPAKAI
            # =================================================

            remaining_mask = ~assigned_mask


            # =================================================
            # HASIL PARAMETER SAAT INI
            # =================================================
            current_group_mask = (
                group_mask
                & remaining_mask
            )
            
            comment_series.loc[current_group_mask] = group.get("comment", "")
            
            # =================================================
            # TANDAI DATA SUDAH MASUK PARAMETER
            # =================================================
           
            assigned_mask = (
                assigned_mask
                | current_group_mask
            )

            if payload.get("export") is True:

                filter_export_progress["progress"] = min(
                    25 + int(
                        (group_index / max(len(filters), 1)) * 45
                    ),
                    70
                )

                filter_export_progress["status"] = (
                    f"Memproses parameter "
                    f"{group_index} / {len(filters)}..."
                )

        # =================================================
        # MASK AKHIR
        # =================================================

        mask = assigned_mask

        # =================================================
        # HASIL FILTER
        # =================================================

        result_df = df.loc[
            mask
        ].copy()

        comment_column_title = payload.get("commentColumnTitle", "Keterangan Penataan")
        result_df[comment_column_title] = comment_series.loc[mask].values


        total = len(
            result_df
        )
                
        export_mode = payload.get(
            "export_mode",
            "filtered"
        )

        if export_mode == "full":
            export_df = df.copy()

            export_df[comment_column_title] = \
                comment_series.values

        else:
            export_df = result_df

        # =================================================
        # EXPORT KE EXCEL
        # =================================================

        if payload.get("export") is True:
            
            filter_export_progress["progress"] = 75
            filter_export_progress["status"] = (
                "Menyiapkan file Excel..."
            )
            
            output = BytesIO()
            
            export_df.to_excel(
                output,
                index=False,
                sheet_name="Hasil Filter",
                engine="openpyxl"
            )

            filter_export_progress["progress"] = 95
            filter_export_progress["status"] = (
                "File Excel berhasil dibuat..."
            )

            output.seek(0)

            safe_sheet_name = "".join(
                c
                if c.isalnum() or c in (" ", "-", "_")
                else "_"
                for c in sheet_name
            ).strip()

            filename = (
                f"hasil_filter_{safe_sheet_name}.xlsx"
            )

            filter_export_progress["progress"] = 100
            filter_export_progress["status"] = "Export selesai."
            filter_export_progress["done"] = True

            return send_file(
                output,
                as_attachment=True,
                download_name=filename,
                mimetype=(
                    "application/"
                    "vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                )
            )

        # -------------------------------------------------
        # TENTUKAN DATA YANG DITAMPILKAN
        # -------------------------------------------------

        if limit is None:

            preview_df = result_df

        else:

            preview_df = result_df.head(
                limit
            )


        # -------------------------------------------------
        # CONVERT KE JSON
        # -------------------------------------------------

        rows = preview_df.to_dict(
            orient="records"
        )


        rows = clean_json_value(
            rows
        )


        # -------------------------------------------------
        # RESPONSE
        # -------------------------------------------------

        return jsonify({
            "success": True,
            "columns": [
                str(column)
                for column in result_df.columns
            ],
            "rows": rows,
            "total": total,
            "shown": len(
                preview_df
            ),
            "limit": limit,
            "show_all": limit is None
        })


    except Exception as error:

        print(
            "[FILTER ERROR]",
            repr(error)
        )

        return jsonify({
            "success": False,
            "message": str(error)
        }), 500

@app.route("/api/header-values")
def header_values():

    filepath = get_current_excel()

    sheet_name = request.args.get(
        "sheet"
    )

    header_row = request.args.get(
        "header",
        type=int
    )

    if filepath is None:
        return jsonify({
            "success": False,
            "message": "Belum ada file Excel."
        }), 404

    if not sheet_name or not header_row:
        return jsonify({
            "success": False,
            "message": "Sheet dan header wajib dipilih."
        }), 400

    try:

        headers = get_header_values(
            filepath,
            sheet_name,
            header_row
        )

        return jsonify({
            "success": True,
            "headers": headers
        })

    except Exception as error:

        return jsonify({
            "success": False,
            "message": str(error)
        }), 500



@app.route("/api/normalize", methods=["POST"])
def normalize():

    payload = request.get_json(
        silent=True
    ) or {}

    value = payload.get(
        "value",
        ""
    )

    parameter = payload.get(
        "parameter",
        ""
    )

    if parameter == "varietas":

        result = normalize_variety(
            value
        )

    else:

        result = normalize_value(
            value
        )

    return jsonify({
        "success": True,
        "original": value,
        "normalized": result,
        "parameter": parameter
    })

# =====================================================
# SHAPEFILE EDITOR
# =====================================================

@app.route("/api/shapefile/upload", methods=["POST"])
def upload_shapefile():

    if "file" not in request.files:
        return jsonify({
            "success": False,
            "message": "File Shapefile tidak ditemukan."
        }), 400

    file = request.files["file"]

    if not file.filename:
        return jsonify({
            "success": False,
            "message": "Silakan pilih file Shapefile ZIP."
        }), 400

    extension = Path(
        file.filename
    ).suffix.lower()

    if extension not in {".zip", ".rar"}:
        return jsonify({
            "success": False,
            "message": (
                "Shapefile harus dikirim "
                "dalam bentuk .zip atau .rar."
            )
        }), 400

    # -------------------------------------------------
    # BERSIHKAN FOLDER SHAPEFILE
    # -------------------------------------------------

    for item in SHAPE_DIR.iterdir():

        if item.is_dir():

            import shutil

            shutil.rmtree(item)

        elif item.is_file():

            item.unlink()

    # -------------------------------------------------
    # SIMPAN ZIP
    # -------------------------------------------------

    zip_path = SHAPE_DIR / file.filename

    file.save(zip_path)

    try:

        extract_dir = (
            SHAPE_DIR / "current"
        )

        shp_path = extract_shapefile(
            zip_path,
            extract_dir
        )

        gdf = read_shapefile(
            shp_path
        )

        geojson = shapefile_to_geojson(
            gdf
        )

        return jsonify({
            "success": True,
            "filename": file.filename,
            "shapefile": shp_path.name,
            "feature_count": len(gdf),
            "columns": [
                str(column)
                for column in gdf.columns
                if column != "geometry"
            ],
            "geojson": geojson
        })

    except Exception as error:

        return jsonify({
            "success": False,
            "message": str(error)
        }), 500

@app.route("/api/shapefile/data")
def shapefile_data():

    try:

        extract_dir = (
            SHAPE_DIR / "current"
        )

        shp_files = list(
            extract_dir.rglob("*.shp")
        )

        if not shp_files:

            return jsonify({
                "success": False,
                "message": (
                    "Belum ada Shapefile "
                    "yang diupload."
                )
            }), 404

        shp_path = shp_files[0]

        gdf = read_shapefile(
            shp_path
        )

        geojson = shapefile_to_geojson(
            gdf
        )

        return jsonify({
            "success": True,
            "filename": shp_path.name,
            "feature_count": len(gdf),
            "columns": [
                str(column)
                for column in gdf.columns
                if column != "geometry"
            ],
            "geojson": geojson
        })

    except Exception as error:

        return jsonify({
            "success": False,
            "message": str(error)
        }), 500
    
@app.route("/api/match/analyze", methods=["POST"])
def analyze_matching():

    try:

        data = request.get_json(
            silent=True
        ) or {}

        # =========================================
        # REQUEST
        # =========================================

        geojson = data.get("geojson")

        excel_sheet = data.get(
            "excel_sheet"
        )

        excel_header = data.get(
            "excel_header"
        )

        shp_key = data.get(
            "shp_key"
        )

        excel_key = data.get(
            "excel_key"
        )

        normalization = data.get(
            "normalization",
            "default"
        )

        # =========================================
        # VALIDASI
        # =========================================

        if not geojson:
            return jsonify({
                "success": False,
                "message": "Data SHP tidak ditemukan."
            }), 400

        if not excel_sheet:
            return jsonify({
                "success": False,
                "message": "Sheet Excel belum dipilih."
            }), 400

        if not excel_header:
            return jsonify({
                "success": False,
                "message": "Header Excel belum dipilih."
            }), 400

        if not shp_key:
            return jsonify({
                "success": False,
                "message": "Atribut kunci SHP belum dipilih."
            }), 400

        if not excel_key:
            return jsonify({
                "success": False,
                "message": "Atribut kunci Excel belum dipilih."
            }), 400

        # =========================================
        # FILE EXCEL
        # =========================================

        filepath = get_current_excel()

        if filepath is None:
            return jsonify({
                "success": False,
                "message": "Belum ada file Excel."
            }), 404

        # =========================================
        # BACA EXCEL
        # =========================================

        excel_df = read_sheet_with_header(
            filepath,
            excel_sheet,
            int(excel_header)
        )

        if excel_df is None:
            return jsonify({
                "success": False,
                "message": "Data Excel tidak dapat dibaca."
            }), 400

        # =========================================
        # SHP GEOJSON → DATAFRAME
        # =========================================

        features = geojson.get(
            "features",
            []
        )

        if not features:
            return jsonify({
                "success": False,
                "message": "Tidak ada polygon SHP."
            }), 400

        shp_rows = []

        for index, feature in enumerate(
            features
        ):

            properties = (
                feature.get("properties") or {}
            )

            row = dict(
                properties
            )

            row["_shape_index"] = index

            shp_rows.append(
                row
            )

        shp_df = pd.DataFrame(
            shp_rows
        )

        # =========================================
        # MATCHING
        # =========================================

        matching_result = match_data(
            shp_df=shp_df,
            excel_df=excel_df,
            shp_key=shp_key,
            excel_key=excel_key,
            normalization=normalization
        )

        # =========================================
        # COMPARE ATTRIBUTE
        # =========================================

        attribute_result = (
            compare_matched_attributes(
                shp_df=shp_df,
                excel_df=excel_df,
                matched=matching_result["matched"],
                shp_key=shp_key,
                excel_key=excel_key,
                normalization=normalization
            )
        )

        # =========================================
        # RESPONSE
        # =========================================

        return jsonify({
            "success": True,

            "matching": matching_result,

            "attributes": attribute_result,

            "summary": {
                "matched": (
                    matching_result["summary"]["matched"]
                ),
                "excel_only": (
                    matching_result["summary"]["excel_only"]
                ),
                "shp_only": (
                    matching_result["summary"]["shp_only"]
                ),
                "same": (
                    attribute_result["summary"]["same"]
                ),
                "new_attributes": (
                    attribute_result["summary"]["new_attributes"]
                ),
                "conflicts": (
                    attribute_result["summary"]["conflicts"]
                )
            }
        })

    except Exception as error:

        print(
            "[MATCH ANALYZE ERROR]",
            repr(error)
        )

        return jsonify({
            "success": False,
            "message": str(error)
        }), 500

@app.route("/api/export-progress")
def export_progress_status():

    return jsonify({
        "success": True,
        "active": export_progress["active"],
        "current": export_progress["current"],
        "total": export_progress["total"],
        "percent": export_progress["percent"],
        "status": export_progress["status"],
        "message": export_progress["message"]
    })    
        
@app.route("/api/export-shp", methods=["POST"])
def export_shp():
    try:
        export_progress["active"] = True
        export_progress["current"] = 0
        export_progress["total"] = 0
        export_progress["percent"] = 0
        export_progress["status"] = "processing"
        export_progress["message"] = "Menyiapkan data SHP..."
        
        data = request.get_json(silent=True) or {}
        geojson = data.get("geojson")

        if not geojson:
            return jsonify({
                "success": False,
                "message": "GeoJSON tidak ditemukan."
            }), 400

        features = geojson.get("features", [])

        if not features:
            return jsonify({
                "success": False,
                "message": "Tidak ada polygon yang dapat diekspor."
            }), 400

        export_progress["total"] = len(features)
        export_progress["current"] = 0
        export_progress["percent"] = 0
        export_progress["message"] = (
            f"Memproses {len(features):,} polygon..."
        )

        print("======================================")
        print("[EXPORT SHP] Mulai")
        print("[EXPORT SHP] Jumlah feature:", len(features))

        gdf_parts = []

        total_features = len(features)

        batch_size = 500

        for start in range(
            0,
            total_features,
            batch_size
        ):

            end = min(
                start + batch_size,
                total_features
            )

            batch_features = features[
                start:end
            ]

            print(
                f"[EXPORT SHP] Batch "
                f"{start + 1:,} - {end:,}"
            )

            part = gpd.GeoDataFrame.from_features(
                batch_features
            )

            gdf_parts.append(
                part
            )

            current = end

            progress_percent = int(
                (current / total_features) * 70
            )

            if current > export_progress["current"]:
                export_progress["current"] = current

            if progress_percent > export_progress["percent"]:
                export_progress["percent"] = progress_percent

            export_progress["message"] = (
                f"Memproses polygon "
                f"{current:,} / "
                f"{total_features:,}"
            )

        gdf = pd.concat(
            gdf_parts,
            ignore_index=True
        )
                
        export_progress["percent"] = 75
        export_progress["message"] = "Menyiapkan file SHP..."        
       
        print(
            "[EXPORT SHP] GeoDataFrame:",
            len(gdf),
            "baris"
        )

        if gdf.empty:
            return jsonify({
                "success": False,
                "message": "GeoDataFrame kosong."
            }), 400

        if gdf.crs is None:
            gdf = gdf.set_crs("EPSG:4326")

        print(
            "[EXPORT SHP] CRS:",
            gdf.crs
        )

        export_dir = DATA_DIR / "export_shp"
        export_dir.mkdir(
            parents=True,
            exist_ok=True
        )

        shp_path = (
            export_dir /
            "variety_engine_export.shp"
        )

        print(
            "[EXPORT SHP] Menulis:",
            shp_path
        )

        gdf.to_file(
            shp_path,
            driver="ESRI Shapefile",
            encoding="UTF-8"
        )
        
        export_progress["percent"] = 90
        export_progress["message"] = "File SHP berhasil dibuat..."  

        print(
            "[EXPORT SHP] File SHP berhasil dibuat"
        )

        # =========================================
        # BUAT ZIP
        # =========================================

        import zipfile

        zip_path = (
            export_dir /
            "variety_engine_export.zip"
        )

        print(
            "[EXPORT SHP] Membuat ZIP:",
            zip_path
        )

        shapefile_extensions = [
            ".shp",
            ".shx",
            ".dbf",
            ".prj",
            ".cpg"
        ]

        with zipfile.ZipFile(
            zip_path,
            "w",
            zipfile.ZIP_DEFLATED
        ) as zip_file:

            for extension in shapefile_extensions:

                component = (
                    export_dir /
                    f"variety_engine_export{extension}"
                )

                if component.exists():

                    print(
                        "[EXPORT SHP] Tambah:",
                        component.name
                    )

                    zip_file.write(
                        component,
                        arcname=component.name
                    )
                    
                    export_progress["percent"] = min(
                        99,
                        export_progress["percent"] + 2
                    )

                    export_progress["message"] = (
                        f"Membuat ZIP... {component.name}"
                    )

        export_progress["current"] = export_progress["total"]
        export_progress["percent"] = 100
        export_progress["status"] = "completed"
        export_progress["active"] = False
        export_progress["message"] = "Export SHP selesai."

        print(
            "[EXPORT SHP] ZIP berhasil dibuat"
        )

        print(
            "[EXPORT SHP] Mengirim ZIP ke browser"
        )

        return send_file(
            zip_path,
            as_attachment=True,
            download_name="variety_engine_export.zip",
            mimetype="application/zip"
        )

    except Exception as error:

        print(
            "[EXPORT SHP ERROR]",
            repr(error)
        )

        return jsonify({
            "success": False,
            "message": str(error)
        }), 500
        
@app.route("/api/export-excel", methods=["POST"])
def export_excel():
    try:
        export_progress["active"] = True
        export_progress["current"] = 0
        export_progress["total"] = 0
        export_progress["percent"] = 0
        export_progress["status"] = "processing"
        export_progress["message"] = "Menyiapkan data Excel..."
                
        data = request.get_json(silent=True) or {}
        geojson = data.get("geojson")

        if not geojson:
            return jsonify({
                "success": False,
                "message": "GeoJSON tidak ditemukan."
            }), 400

        features = geojson.get("features", [])

        if not features:
            return jsonify({
                "success": False,
                "message": "Tidak ada data yang dapat diekspor."
            }), 400

        export_progress["total"] = len(features)
        export_progress["current"] = 0
        export_progress["percent"] = 0
        export_progress["message"] = (
            f"Menyiapkan {len(features):,} polygon untuk Excel..."
        )

        print("======================================")
        print("[EXPORT EXCEL] Mulai")
        print("[EXPORT EXCEL] Jumlah feature:", len(features))

        # -----------------------------------------
        # AMBIL ATRIBUT DARI GEOJSON
        # -----------------------------------------

        rows = []

        total_features = len(features)

        for index, feature in enumerate(features, start=1):

            properties = (
                feature.get("properties") or {}
            )

            rows.append(properties)

            progress_percent = int(
                (index / total_features) * 70
            )

            if progress_percent > export_progress["percent"]:
                export_progress["percent"] = progress_percent

            export_progress["current"] = index

            export_progress["message"] = (
                f"Menyiapkan data "
                f"{index:,} / {total_features:,} polygon..."
            )

        print(
            "[EXPORT EXCEL] Jumlah baris:",
            len(rows)
        )

        # -----------------------------------------
        # BUAT DATAFRAME
        # -----------------------------------------

        df = pd.DataFrame(rows)

        export_progress["percent"] = 80
        export_progress["message"] = "Menyiapkan struktur file Excel..."

        print(
            "[EXPORT EXCEL] Jumlah kolom:",
            len(df.columns)
        )

        # -----------------------------------------
        # BUAT FILE EXCEL DI MEMORY
        # -----------------------------------------

        output = BytesIO()

        export_progress["percent"] = 85
        export_progress["message"] = "Membuat file Excel..."

        output = BytesIO()

        export_progress["percent"] = 50
        export_progress["message"] = "Mulai membuat file Excel..."

        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Data Atribut"

        # -----------------------------------------
        # TULIS HEADER
        # -----------------------------------------

        for column_index, column_name in enumerate(
            df.columns,
            start=1
        ):
            worksheet.cell(
                row=1,
                column=column_index,
                value=column_name
            )

        # -----------------------------------------
        # TULIS DATA + PROGRESS
        # -----------------------------------------

        total_rows = len(df)

        for row_index, row in enumerate(
            df.itertuples(index=False, name=None),
            start=2
        ):

            for column_index, value in enumerate(
                row,
                start=1
            ):
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=value
                )

            current_row = row_index - 1

            progress_percent = (
                50
                + int(
                    (current_row / total_rows) * 49
                )
            )

            export_progress["current"] = current_row
            export_progress["percent"] = min(
                progress_percent,
                99
            )

            export_progress["message"] = (
                f"Membuat file Excel... "
                f"{current_row:,} / "
                f"{total_rows:,} baris"
            )

        # -----------------------------------------
        # SIMPAN FILE EXCEL
        # -----------------------------------------

        workbook.save(output)

        export_progress["current"] = total_rows
        export_progress["percent"] = 100
        export_progress["status"] = "completed"
        export_progress["active"] = False
        export_progress["message"] = "Export Excel selesai."

        output.seek(0)

        print(
            "[EXPORT EXCEL] File Excel berhasil dibuat"
        )

        print(
            "[EXPORT EXCEL] Mengirim Excel ke browser"
        )

        return send_file(
            output,
            as_attachment=True,
            download_name="variety_engine_export.xlsx",
            mimetype=(
                "application/"
                "vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

    except Exception as error:

        print(
            "[EXPORT EXCEL ERROR]",
            repr(error)
        )

        return jsonify({
            "success": False,
            "message": str(error)
        }), 500

if __name__ == "__main__":

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=True
    )
